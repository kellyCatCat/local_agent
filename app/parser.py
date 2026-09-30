"""把上传的 Excel / Markdown / CSV / 文本转换成给模型阅读的 markdown 文本。"""
from __future__ import annotations

import csv
import io
from pathlib import PurePath

from openpyxl import load_workbook

TEXT_EXTS = {".md", ".markdown", ".txt"}
EXCEL_EXTS = {".xlsx", ".xlsm"}
CSV_EXTS = {".csv"}
SUPPORTED_EXTS = TEXT_EXTS | EXCEL_EXTS | CSV_EXTS


class ParseError(ValueError):
    pass


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _cell(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return s.replace("\r\n", "\n").replace("\r", "\n").replace("|", "\\|").replace("\n", "<br>")


def rows_to_markdown(rows: list[list]) -> str:
    """二维数组 -> markdown 表格；去掉全空行/全空列，第一行作表头。"""
    rows = [[_cell(v) for v in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return "（空表）"
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [i for i in range(width) if any(r[i] for r in rows)]
    rows = [[r[i] for i in keep] for r in rows]
    header, body = rows[0], rows[1:]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join("---" for _ in header) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def _excel_grids(data: bytes) -> list[tuple[str, list[list]]]:
    """读取每个 sheet 的二维数据；合并单元格把左上角的值填充到整个合并区域。"""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:  # openpyxl 抛出的异常类型较多
        raise ParseError(f"无法解析 Excel：{e}") from e
    out = []
    for ws in wb.worksheets:
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        for rng in ws.merged_cells.ranges:
            v = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                        grid[r - 1][c - 1] = v
        if any(any(v is not None and str(v).strip() for v in row) for row in grid):
            out.append((ws.title, grid))
    return out


def excel_to_markdown(data: bytes) -> str:
    parts = [f"### Sheet：{title}\n\n{rows_to_markdown(grid)}" for title, grid in _excel_grids(data)]
    return "\n\n".join(parts) if parts else "（Excel 中没有内容）"


# 模板规定的命令来源列；「回显」列不是命令来源
CLI_SOURCE_HEADERS = ("命令行", "配置修复建议", "修复验证", "步骤详细描述")
ECHO_HEADER = "回显"
HEADER_SCAN_ROWS = 10


def excel_cli_source(data: bytes) -> dict | None:
    """按表头提取步骤表的命令来源列与回显列文本，供 CLI 来源检查使用。

    表头单元格包含「命令行」「配置修复建议」「修复验证」「步骤详细描述」之一即视为命令来源列
    （如「修复验证，怎么验证」）。没有任何 sheet 能识别出命令来源列时返回 None。
    """
    cmd, echo, sheets = [], [], []
    for title, grid in _excel_grids(data):
        for hi, row in enumerate(grid[:HEADER_SCAN_ROWS]):
            heads = [str(v or "").strip() for v in row]
            src_cols = [i for i, h in enumerate(heads) if any(k in h for k in CLI_SOURCE_HEADERS)]
            if not src_cols:
                continue
            echo_cols = [i for i, h in enumerate(heads) if ECHO_HEADER in h and i not in src_cols]
            for r in grid[hi + 1:]:
                cmd += [str(r[i]) for i in src_cols if i < len(r) and r[i] is not None]
                echo += [str(r[i]) for i in echo_cols if i < len(r) and r[i] is not None]
            sheets.append(title)
            break
    if not sheets:
        return None
    return {"cmd": "\n".join(cmd), "echo": "\n".join(echo), "sheets": sheets}


def parse_cli_source(filename: str, data: bytes) -> dict | None:
    """目前只有 Excel 步骤表能区分命令来源列与回显列；其他格式返回 None（不做 CLI 来源检查）。"""
    if PurePath(filename).suffix.lower() in EXCEL_EXTS:
        return excel_cli_source(data)
    return None


def parse_upload(filename: str, data: bytes) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext in TEXT_EXTS:
        return _decode(data)
    if ext in EXCEL_EXTS:
        return excel_to_markdown(data)
    if ext in CSV_EXTS:
        return rows_to_markdown(list(csv.reader(io.StringIO(_decode(data)))))
    if ext == ".xls":
        raise ParseError("不支持旧版 .xls，请在 Excel 中另存为 .xlsx 后上传")
    raise ParseError(f"不支持的文件类型：{ext or filename}（支持 {', '.join(sorted(SUPPORTED_EXTS))}）")

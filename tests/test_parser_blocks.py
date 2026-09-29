import io

import pytest
from openpyxl import Workbook

from app.fileblocks import UnsafePathError, parse_output, safe_relpath, strip_blocks
from app.parser import ParseError, parse_upload


def test_excel_merged_cells_and_escaping():
    wb = Workbook()
    ws = wb.active
    ws.title = "步骤表"
    ws.append(["子场景", "步骤详细描述", "命令行", "回显"])
    ws.append(["邻居Down", "检查接口", "display interface <端口>", "a|b\nc"])
    ws.append([None, "检查邻居", "display isis peer", None])
    ws.merge_cells("A2:A3")
    wb.create_sheet("空表")
    buf = io.BytesIO()
    wb.save(buf)
    md = parse_upload("steps.xlsx", buf.getvalue())
    assert "### Sheet：步骤表" in md
    assert "空表" not in md
    assert md.count("邻居Down") == 2  # 合并单元格被填充
    assert "a\\|b<br>c" in md
    assert "| 子场景 | 步骤详细描述 | 命令行 | 回显 |" in md


def test_text_and_unsupported():
    assert parse_upload("a.md", "# 标题".encode("gb18030")) == "# 标题"
    assert "| x | y |" in parse_upload("a.csv", b"x,y\n1,2\n")
    with pytest.raises(ParseError):
        parse_upload("a.xls", b"")
    with pytest.raises(ParseError):
        parse_upload("a.pdf", b"")


def test_parse_output_blocks():
    text = (
        "改动说明\n<<<FILE: SKILL.md>>>\n---\nname: x\n---\n```\ncode\n```\n<<<END FILE>>>\n"
        "<<<FILE: reference/a.md>>>\n```markdown\n# 场景A：x\n```\n<<<END FILE>>>\n"
        "<<<DELETE: reference/old.md>>>\n"
        "<<<FILE: ../evil.md>>>\nx\n<<<END FILE>>>\n"
        "<<<FILE: reference/cut.md>>>\n半截"
    )
    out = parse_output(text)
    assert out.files["SKILL.md"] == "---\nname: x\n---\n```\ncode\n```\n"
    assert out.files["reference/a.md"] == "# 场景A：x\n"  # 外层围栏被剥离
    assert out.deletes == ["reference/old.md"]
    assert out.truncated == ["reference/cut.md"]
    assert out.errors and "../evil.md" in out.errors[0]
    assert "[已输出文件 SKILL.md]" in strip_blocks(text)


@pytest.mark.parametrize("bad", ["/etc/passwd", "../x", "a/../../x", "", "C:/x", "~/x"])
def test_safe_relpath_rejects(bad):
    with pytest.raises(UnsafePathError):
        safe_relpath(bad)


def test_safe_relpath_normalizes():
    assert safe_relpath("./reference\\a.md") == "reference/a.md"
